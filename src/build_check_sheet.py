"""Build the human-check workbook: all needs_review cells + a random sample of pass cells."""
import json, random, glob, pathlib, sys
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from llm_client import ROOT
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

VARIANTS = {"hi_deva": "Hindi (Devanagari)", "hi_roman": "Romanized Hindi", "hi_cm": "Hinglish (code-mixed)"}
DROPPED = {"Natural-007", "GPTInst-079"}  # translator safety-blocked; excluded item-level
SAMPLE_PER_VARIANT = 30
en = {json.loads(l)["id"]: json.loads(l) for l in open(ROOT / "data" / "llmbar_en.jsonl")}

def find_issues(o):
    if isinstance(o, dict):
        if isinstance(o.get("issues"), str): return o["issues"]
        for v in o.values():
            r = find_issues(v)
            if r: return r
    return ""

issues = {}
for v in VARIANTS:
    for l in open(ROOT / "data" / "raw" / "verification_attempts" / f"{v}.jsonl"):
        d = json.loads(l); issues[(d["id"], v)] = find_issues(d)  # last record wins (latest repair round)

wb = Workbook(); ws0 = wb.active; ws0.title = "README"
for row in [
    ["HUMAN CHECK - how to fill"],
    ["Each person fills their OWN copy. Do not discuss answers until both are done."],
    ["Type = REVIEW: the automatic checker found a possible problem (see 'Checker notes'). Type = SAMPLE: passed all checks; checked for the agreement score."],
    ["Meaning: does the converted text mean the same as the English? Yes / Partly / No (judge instruction + both outputs together)."],
    ["Label: is the 'Better output' still the better one, for the same reasons, in the converted version? Yes / No"],
    ["Naturalness: does it read like real text a person would write in this form? 1 (robotic) - 5 (natural)"],
    ["Fix (optional): if you can fix the translation, paste the corrected text in 'Suggested fix' and say which field (instruction/output_1/output_2)."],
]: ws0.append(row)
ws0.column_dimensions["A"].width = 140

rng = random.Random(7)
heads = ["#", "Type", "Item id", "Better output", "Checker notes",
         "Instruction (EN)", "Instruction (converted)", "Output 1 (EN)", "Output 1 (converted)",
         "Output 2 (EN)", "Output 2 (converted)",
         "Meaning (Yes/Partly/No)", "Label (Yes/No)", "Naturalness (1-5)", "Suggested fix", "Notes"]
widths = [5, 9, 15, 9, 40, 40, 40, 50, 50, 50, 50, 14, 11, 12, 40, 30]
counts = {}
for v, title in VARIANTS.items():
    recs = [json.loads(l) for l in open(ROOT / "data" / "translations" / f"{v}.jsonl")]
    recs = [r for r in recs if r["id"] not in DROPPED]
    rev = sorted([r for r in recs if r["status"] == "needs_review"], key=lambda r: r["id"])
    pas = [r for r in recs if r["status"] == "pass"]
    samp = sorted(rng.sample(pas, SAMPLE_PER_VARIANT), key=lambda r: r["id"])
    ws = wb.create_sheet(title[:31]); ws.append(heads)
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEBF7")
    n = 0
    for typ, group in (("REVIEW", rev), ("SAMPLE", samp)):
        for r in group:
            n += 1; e = en[r["id"]]; t = r["texts"]
            note = issues.get((r["id"], v), "") or "; ".join(r.get("reasons", []))
            ws.append([n, typ, r["id"], f"Output {e['label']}", note,
                       e["input"], t.get("input", ""), e["output_1"], t.get("output_1", ""),
                       e["output_2"], t.get("output_2", ""), "", "", "", "", ""])
    for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "F2"
    for col, opts in (("L", '"Yes,Partly,No"'), ("M", '"Yes,No"'), ("N", '"1,2,3,4,5"')):
        dv = DataValidation(type="list", formula1=opts, allow_blank=True); ws.add_data_validation(dv); dv.add(f"{col}2:{col}{n+1}")
    counts[title] = (len(rev), len(samp))
wb.save(ROOT / "data" / "human_validation" / "blank_sheet.xlsx")
print(counts)
